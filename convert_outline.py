# -*- coding: utf-8 -*-
"""讲解大纲.html -> 讲解大纲.docx （原生 Word 标题/表格/列表/代码块，便于复制写培训记录）"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import html as lhtml

SRC = "讲解大纲.html"
OUT = "讲解大纲.docx"
CN_FONT = "微软雅黑"

TAG_COLORS = {
    "tag-blue": "2563EB", "tag-purple": "6366F1", "tag-green": "059669",
    "tag-orange": "B45309", "tag-red": "DC2626",
}
SKIP_TAGS = {"nav", "script", "style", "link", "meta", "head", "title"}

doc = Document()


# ---------- helpers ----------
def set_east_asia(style, font_name):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


# ---------- fonts ----------
normal = doc.styles["Normal"]
normal.font.name = CN_FONT
normal.font.size = Pt(10.5)
set_east_asia(normal, CN_FONT)
for hs in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
    try:
        set_east_asia(doc.styles[hs], CN_FONT)
    except KeyError:
        pass


# ---------- inline ----------
def render_inline(el, p, bold=False, mono=False, color=None):
    def add_text(t):
        if not t:
            return
        run = p.add_run(t)
        if bold:
            run.bold = True
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    add_text(el.text or "")
    for c in el:
        if not isinstance(c.tag, str):
            add_text(c.tail or "")
            continue
        tag = c.tag
        if tag in ("b", "strong"):
            render_inline(c, p, bold=True, mono=mono, color=color)
        elif tag == "code":
            render_inline(c, p, bold=bold, mono=True, color=color)
        elif tag == "br":
            p.add_run().add_break()
        elif tag == "span":
            tcolor = None
            for cls in (c.get("class") or "").split():
                if cls in TAG_COLORS:
                    tcolor = TAG_COLORS[cls]
            render_inline(c, p, bold=bold or bool(tcolor), mono=mono, color=tcolor or color)
        elif tag in ("ul", "ol", "table", "pre", "div"):
            # block-level handled elsewhere
            pass
        else:
            render_inline(c, p, bold=bold, mono=mono, color=color)
        add_text(c.tail or "")


def text_skip_badge(e):
    out = e.text or ""
    for c in e:
        if not isinstance(c.tag, str):
            continue
        if c.tag == "span" and "badge" in (c.get("class") or ""):
            out += c.tail or ""
            continue
        out += text_skip_badge(c)
        out += c.tail or ""
    return out


def heading_text(e):
    return " ".join(text_skip_badge(e).split())


# ---------- block ----------
def render_list(el, ordered, depth=0):
    n = 0
    for li in el:
        if not isinstance(li.tag, str) or li.tag != "li":
            continue
        n += 1
        if ordered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5 + 0.3 * depth)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            num = p.add_run(f"{n}. ")
            num.bold = True
            render_inline(li, p)
        else:
            name = ["List Bullet", "List Bullet 2", "List Bullet 3"][min(depth, 2)]
            try:
                doc.styles[name]
            except KeyError:
                name = "List Bullet"
            p = doc.add_paragraph(style=name)
            if depth and name == "List Bullet":
                p.paragraph_format.left_indent = Inches(0.75 + 0.3 * (depth - 1))
            render_inline(li, p)
        for c in li:
            if isinstance(c.tag, str) and c.tag in ("ul", "ol"):
                render_list(c, ordered=(c.tag == "ol"), depth=depth + 1)


def render_table(el):
    rows = el.xpath(".//tr")
    if not rows:
        return
    ncols = 0
    for r in rows:
        ncols = max(ncols, len([c for c in r if isinstance(c.tag, str) and c.tag in ("td", "th")]))
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    for ri, r in enumerate(rows):
        cells = [c for c in r if isinstance(c.tag, str) and c.tag in ("td", "th")]
        for ci, c in enumerate(cells):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            render_inline(c, p)
            for run in p.runs:
                run.font.size = Pt(10)
            if c.tag == "th":
                shade_cell(cell, "D9DCF7")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def render_pre(el):
    text = el.text_content()
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade_cell(cell, "F4F4F8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        if i < len(lines) - 1:
            run.add_break()
    doc.add_paragraph()


def render_callout(el, kind):
    fill = "EAF2FF" if kind == "note" else "FFF4E5"
    p = doc.add_paragraph()
    shade_paragraph(p, fill)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    render_inline(el, p)


def process_block(el):
    if not isinstance(el.tag, str):
        return
    tag = el.tag
    if tag in SKIP_TAGS:
        return
    if tag == "h2":
        doc.add_heading(heading_text(el), level=1)
    elif tag == "h3":
        doc.add_heading(heading_text(el), level=2)
    elif tag == "h4":
        doc.add_heading(heading_text(el), level=3)
    elif tag == "p":
        render_inline(el, doc.add_paragraph())
    elif tag == "ul":
        render_list(el, ordered=False)
    elif tag == "ol":
        render_list(el, ordered=True)
    elif tag == "table":
        render_table(el)
    elif tag == "pre":
        render_pre(el)
    elif tag == "div":
        cls = el.get("class") or ""
        if "note" in cls:
            render_callout(el, "note")
        elif "warn" in cls:
            render_callout(el, "warn")
        else:
            for c in el:
                process_block(c)


# ---------- run ----------
tree = lhtml.parse(SRC)
body = tree.getroot().find("body")
title = (tree.find(".//title").text or "讲解大纲").strip()
doc.add_heading(title, level=0)
for child in body:
    process_block(child)

doc.save(OUT)
print(f"OK -> {OUT}")

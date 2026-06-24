# -*- coding: utf-8 -*-
"""从 config.json 读 Excel 路径，查找 CK-FCM-2026039。"""
import json, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

cfg = json.load(open('config.json', encoding='utf-8'))
path = cfg['defaults']['excelPath']
print('path:', path)
import os
print('exists:', os.path.exists(path))

wb = openpyxl.load_workbook(path, read_only=True)
print('sheets:', wb.sheetnames)
TARGET = 'CK-FCM-2026039'
for sn in wb.sheetnames:
    ws = wb[sn]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: continue
    hdr = rows[0]
    for i, row in enumerate(rows[1:], start=2):
        cells = [str(c) if c is not None else '' for c in row]
        if any(TARGET in c for c in cells):
            print(f'\n=== HIT in sheet [{sn}] row {i} ===')
            for h, v in zip(hdr, row):
                print(f'  {h} = {v!r}')
